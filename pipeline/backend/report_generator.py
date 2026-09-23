import io
import zipfile
from datetime import datetime, timezone

import openpyxl
from fpdf import FPDF

import github_integration


_FAMILIARITY_LABELS = {
    "first_time": "First time",
    "used_before": "Used it before",
    "regular": "Regular user",
}
_COMFORT_LABELS = {
    "basic": "Basic",
    "comfortable": "Comfortable",
    "advanced": "Advanced",
}


def _safe(text: str) -> str:
    """Encode to latin-1, replacing unsupported characters, so fpdf2 doesn't crash."""
    return text.encode("latin-1", errors="replace").decode("latin-1")


def _parse_reporter(body: str) -> dict:
    """Extract the ## Reporter table from an issue body into a flat dict."""
    start = body.find("## Reporter")
    if start == -1:
        return {}
    end = body.find("\n## ", start + 1)
    section = body[start:end] if end != -1 else body[start:]
    result: dict[str, str] = {}
    for line in section.splitlines():
        if not line.startswith("|"):
            continue
        parts = [p.strip() for p in line.split("|") if p.strip()]
        if len(parts) == 2 and not all(c in "-" for c in parts[0]):
            key = parts[0].lower()
            if key != "field":
                result[key] = parts[1]
    return result


def _extract_summary(body: str) -> str:
    start = body.find("## Summary")
    if start == -1:
        return ""
    text_start = body.find("\n", start) + 1
    end = body.find("\n## ", text_start)
    block = body[text_start:end] if end != -1 else body[text_start:]
    return block.strip()


def _label_value(labels: list[dict], prefix: str) -> str:
    for lbl in labels:
        name = lbl.get("name", "")
        if name.startswith(f"{prefix}::"):
            return name.split("::", 1)[1]
    return ""


def _parse_issues(raw_issues: list[dict]) -> list[dict]:
    parsed = []
    for issue in raw_issues:
        body = issue.get("body") or ""
        labels = issue.get("labels", [])
        reporter = _parse_reporter(body)
        parsed.append({
            "number": issue["number"],
            "title": issue["title"],
            "url": issue["html_url"],
            "type": _label_value(labels, "type"),
            "scope": _label_value(labels, "scope"),
            "priority": _label_value(labels, "priority"),
            "reporter": reporter,
            "summary": _extract_summary(body),
            "created_at": issue.get("created_at", ""),
        })
    return parsed


def _unique_participants(issues: list[dict]) -> list[dict]:
    seen: set[str] = set()
    participants = []
    for issue in issues:
        r = issue["reporter"]
        name = r.get("name", "")
        if name and name not in seen:
            seen.add(name)
            participants.append(r)
    return participants


# ── Markdown ──────────────────────────────────────────────────────────────────

def _generate_markdown(tag: str, issues: list[dict]) -> str:
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    participants = _unique_participants(issues)

    type_counts: dict[str, int] = {}
    priority_counts: dict[str, int] = {}
    scope_counts: dict[str, int] = {}
    for issue in issues:
        type_counts[issue["type"] or "unknown"] = type_counts.get(issue["type"] or "unknown", 0) + 1
        priority_counts[issue["priority"] or "unknown"] = priority_counts.get(issue["priority"] or "unknown", 0) + 1
        scope_counts[issue["scope"] or "unknown"] = scope_counts.get(issue["scope"] or "unknown", 0) + 1

    lines = [
        f"# Workshop Report — {tag}",
        f"_Generated: {now}_  ",
        f"_Total issues: {len(issues)} | Participants: {len(participants)}_",
        "",
        "---",
        "",
        "## Participants",
        "",
        "| Name | Role | Organization | App Familiarity | Experience | Digital Comfort |",
        "|------|------|--------------|-----------------|------------|----------------|",
    ]
    for p in participants:
        familiarity = _FAMILIARITY_LABELS.get(p.get("app familiarity", ""), p.get("app familiarity", ""))
        comfort = _COMFORT_LABELS.get(p.get("digital comfort", ""), p.get("digital comfort", ""))
        lines.append(
            f"| {p.get('name','')} | {p.get('role','')} | {p.get('organization','')} | "
            f"{familiarity} | {p.get('years of experience','')} | {comfort} |"
        )
    if not participants:
        lines.append("| — | — | — | — | — | — |")

    lines += [
        "",
        "---",
        "",
        "## Statistics",
        "",
        "### By Type",
    ]
    for k, v in sorted(type_counts.items()):
        lines.append(f"- **{k}**: {v}")
    lines += ["", "### By Priority"]
    for k, v in sorted(priority_counts.items()):
        lines.append(f"- **{k}**: {v}")
    lines += ["", "### By Scope"]
    for k, v in sorted(scope_counts.items()):
        lines.append(f"- **{k}**: {v}")

    lines += ["", "---", "", "## Issues", ""]
    for issue in issues:
        r = issue["reporter"]
        reporter_name = r.get("name", "anonymous")
        summary = (issue["summary"][:300] + "…") if len(issue["summary"]) > 300 else issue["summary"]
        lines += [
            f"### #{issue['number']} — {issue['title']}",
            f"**Type**: `{issue['type']}` | **Priority**: `{issue['priority']}` | **Scope**: `{issue['scope']}`  ",
            f"**Reporter**: {reporter_name} ({r.get('role','')}, {r.get('organization','')})  ",
            f"**URL**: {issue['url']}  ",
            f"**Created**: {issue['created_at'][:10] if issue['created_at'] else ''}",
            "",
            summary or "_No summary available._",
            "",
        ]

    return "\n".join(lines)


# ── Excel ─────────────────────────────────────────────────────────────────────

def _generate_excel(tag: str, issues: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    participants = _unique_participants(issues)

    # Sheet 1: Participants
    ws_p = wb.active
    ws_p.title = "Participants"
    ws_p.append(["Name", "Role", "Organization", "App Familiarity", "Years Experience", "Digital Comfort"])
    for p in participants:
        familiarity = _FAMILIARITY_LABELS.get(p.get("app familiarity", ""), p.get("app familiarity", ""))
        comfort = _COMFORT_LABELS.get(p.get("digital comfort", ""), p.get("digital comfort", ""))
        ws_p.append([
            p.get("name", ""),
            p.get("role", ""),
            p.get("organization", ""),
            familiarity,
            p.get("years of experience", ""),
            comfort,
        ])

    # Sheet 2: Issues
    ws_i = wb.create_sheet("Issues")
    ws_i.append(["#", "Title", "URL", "Type", "Scope", "Priority", "Reporter", "Role", "Organization", "Created"])
    for issue in issues:
        r = issue["reporter"]
        ws_i.append([
            issue["number"],
            issue["title"],
            issue["url"],
            issue["type"],
            issue["scope"],
            issue["priority"],
            r.get("name", ""),
            r.get("role", ""),
            r.get("organization", ""),
            issue["created_at"][:10] if issue["created_at"] else "",
        ])

    # Sheet 3: Statistics
    ws_s = wb.create_sheet("Statistics")
    ws_s.append(["Metric", "Value"])
    ws_s.append(["Total Issues", len(issues)])
    ws_s.append(["Total Participants", len(participants)])

    type_counts: dict[str, int] = {}
    priority_counts: dict[str, int] = {}
    scope_counts: dict[str, int] = {}
    for issue in issues:
        type_counts[issue["type"] or "unknown"] = type_counts.get(issue["type"] or "unknown", 0) + 1
        priority_counts[issue["priority"] or "unknown"] = priority_counts.get(issue["priority"] or "unknown", 0) + 1
        scope_counts[issue["scope"] or "unknown"] = scope_counts.get(issue["scope"] or "unknown", 0) + 1

    ws_s.append([])
    ws_s.append(["--- By Type ---", ""])
    for k, v in sorted(type_counts.items()):
        ws_s.append([k, v])
    ws_s.append([])
    ws_s.append(["--- By Priority ---", ""])
    for k, v in sorted(priority_counts.items()):
        ws_s.append([k, v])
    ws_s.append([])
    ws_s.append(["--- By Scope ---", ""])
    for k, v in sorted(scope_counts.items()):
        ws_s.append([k, v])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PDF ───────────────────────────────────────────────────────────────────────

def _generate_pdf(tag: str, issues: list[dict]) -> bytes:
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()

    # Title
    pdf.set_font("Helvetica", "B", 20)
    pdf.cell(0, 12, "Workshop Report", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_font("Helvetica", "I", 13)
    pdf.cell(0, 9, _safe(tag), new_x="LMARGIN", new_y="NEXT", align="C")
    now = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(100, 100, 100)
    pdf.cell(0, 6, f"Generated: {now}", new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.set_text_color(0, 0, 0)
    pdf.ln(8)

    participants = _unique_participants(issues)

    # Statistics summary
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Overview", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 7, f"Total issues: {len(issues)}", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, f"Participants: {len(participants)}", new_x="LMARGIN", new_y="NEXT")

    type_counts: dict[str, int] = {}
    priority_counts: dict[str, int] = {}
    for issue in issues:
        t = issue["type"] or "unknown"
        p = issue["priority"] or "unknown"
        type_counts[t] = type_counts.get(t, 0) + 1
        priority_counts[p] = priority_counts.get(p, 0) + 1

    type_str = ", ".join(f"{k}: {v}" for k, v in sorted(type_counts.items()))
    priority_str = ", ".join(f"{k}: {v}" for k, v in sorted(priority_counts.items()))
    pdf.cell(0, 7, _safe(f"By type — {type_str}"), new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 7, _safe(f"By priority — {priority_str}"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(6)

    # Participants table
    if participants:
        pdf.set_font("Helvetica", "B", 14)
        pdf.cell(0, 10, "Participants", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "B", 9)
        col_w = [45, 40, 50, 35]
        headers = ["Name", "Role", "Organization", "App Familiarity"]
        for i, h in enumerate(headers):
            pdf.cell(col_w[i], 7, h, border=1)
        pdf.ln()
        pdf.set_font("Helvetica", "", 9)
        for p in participants:
            familiarity = _FAMILIARITY_LABELS.get(p.get("app familiarity", ""), p.get("app familiarity", ""))
            row = [
                _safe(p.get("name", "")),
                _safe(p.get("role", "")),
                _safe(p.get("organization", "")),
                _safe(familiarity),
            ]
            for i, cell in enumerate(row):
                pdf.cell(col_w[i], 7, cell[:22] if len(cell) > 22 else cell, border=1)
            pdf.ln()
        pdf.ln(6)

    # Issues list
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 10, "Issues", new_x="LMARGIN", new_y="NEXT")

    for issue in issues:
        r = issue["reporter"]
        pdf.set_font("Helvetica", "B", 10)
        title = _safe(f"#{issue['number']} — {issue['title']}")
        if len(title) > 90:
            title = title[:87] + "..."
        pdf.multi_cell(0, 7, title)
        pdf.set_font("Helvetica", "", 9)
        meta = _safe(f"Type: {issue['type']} | Priority: {issue['priority']} | Scope: {issue['scope']}")
        pdf.cell(0, 6, meta, new_x="LMARGIN", new_y="NEXT")
        if r.get("name"):
            rep = _safe(f"Reporter: {r.get('name','')} — {r.get('role','')} @ {r.get('organization','')}")
            if len(rep) > 90:
                rep = rep[:87] + "..."
            pdf.cell(0, 6, rep, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(30, 100, 200)
        url = issue["url"]
        if len(url) > 90:
            url = url[:87] + "..."
        pdf.cell(0, 6, url, new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(0, 0, 0)
        if issue["summary"]:
            summary = _safe(issue["summary"][:250])
            if len(issue["summary"]) > 250:
                summary += "..."
            pdf.multi_cell(0, 6, summary)
        pdf.ln(3)

    return bytes(pdf.output())


# ── Public entry point ────────────────────────────────────────────────────────

def generate_report_zip(tag: str) -> bytes:
    raw_issues = github_integration.fetch_workshop_issues(tag)
    issues = _parse_issues(raw_issues)

    md = _generate_markdown(tag, issues)
    xlsx = _generate_excel(tag, issues)
    pdf_bytes = _generate_pdf(tag, issues)

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{tag}-report.md", md.encode("utf-8"))
        zf.writestr(f"{tag}-report.pdf", pdf_bytes)
        zf.writestr(f"{tag}-report.xlsx", xlsx)
    return buf.getvalue()
